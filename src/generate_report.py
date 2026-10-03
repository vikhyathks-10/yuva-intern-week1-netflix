import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

PROJECT_DIR = os.path.dirname(os.path.dirname(__file__))
REPORT_DIR = os.path.join(PROJECT_DIR, 'reports')
VIS_DIR = os.path.join(PROJECT_DIR, 'visualizations')
OUTPUT_DOCX = os.path.join(REPORT_DIR, 'Yuva_Internship_Week1_Report.docx')

os.makedirs(REPORT_DIR, exist_ok=True)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def style_table(table, col_widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Add borders
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
                <w:insideV w:val="none"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr[0].append(borders)
    
    # Format header row
    hdr_cells = table.rows[0].cells
    for cell in hdr_cells:
        set_cell_background(cell, "E50914") # Netflix Red
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(10)
                run.font.name = 'Arial'
                
    # Format data rows
    for row_idx, row in enumerate(table.rows[1:], start=1):
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for cell in row.cells:
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80, left=150, right=150)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9.5)
                    run.font.name = 'Arial'
                    run.font.color.rgb = RGBColor(30, 41, 59)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(229, 9, 20) # Netflix Red
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42) # Slate Dark
    return p

def add_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_image_with_caption(doc, img_path, caption):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        run.add_picture(img_path, width=Inches(5.8))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(12)
        run_cap = p_cap.add_run(f"Figure: {caption}")
        run_cap.font.name = 'Arial'
        run_cap.font.size = Pt(9)
        run_cap.font.italic = True
        run_cap.font.color.rgb = RGBColor(100, 116, 139)

def generate_report():
    doc = docx.Document()

    # Set page margins (1 inch on all sides)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # ================= COVER PAGE =================
    p_top_space = doc.add_paragraph()
    p_top_space.paragraph_format.space_before = Pt(30)
    
    p_org = doc.add_paragraph()
    p_org.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_org = p_org.add_run("YUVA INTERNSHIP PROGRAM — WEEK 1 REPORT")
    run_org.font.name = 'Arial'
    run_org.font.size = Pt(12)
    run_org.font.bold = True
    run_org.font.color.rgb = RGBColor(100, 116, 139)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(20)
    p_title.paragraph_format.space_after = Pt(10)
    run_title = p_title.add_run("Netflix Movies and TV Shows:\nData Acquisition, Cleaning, and Preprocessing")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(229, 9, 20)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(140)
    run_sub = p_sub.add_run("A Comprehensive Data Engineering & Exploratory Analytics Pipeline")
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.line_spacing = 1.3
    
    meta_runs = [
        ("Student Name: ", True), ("Vikhyath Bharadwaj K S\n", False),
        ("Domain: ", True), ("Data Science & Data Engineering\n", False),
        ("Dataset: ", True), ("Netflix Movies and TV Shows (Shivam Bansal / Kaggle)\n", False),
        ("Submission Date: ", True), ("October 3, 2026\n", False),
        ("Repository: ", True), ("yuva-intern-week1-netflix", False)
    ]
    for label, is_bold in meta_runs:
        r = p_meta.add_run(label)
        r.font.name = 'Arial'
        r.font.size = Pt(11)
        r.font.bold = is_bold
        r.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_page_break()

    # ================= 2. ABSTRACT =================
    add_heading_1(doc, "1. Abstract")
    add_paragraph(doc, "This technical report details the end-to-end data acquisition, cleaning, preprocessing, and exploratory data analysis (EDA) conducted for Week 1 of the Yuva Internship program. Utilizing the official Netflix Movies and TV Shows dataset curated by Shivam Bansal, this project establishes a robust, reproducible Python pipeline that transforms uncleaned, missing-value-heavy raw metadata into an analysis-ready dataset.")
    add_paragraph(doc, "The original raw dataset comprises 7,787 entries and 12 features spanning content released between 1925 and 2021. Data quality assessment revealed significant missingness in key attributes, including director names (30.68% missing), cast details (9.22% missing), production countries (6.51% missing), date added (0.13% missing), and ratings (0.09% missing). Rather than discarding incomplete records and reducing sample size, transparent domain-justified imputation techniques were employed, maintaining 100% record retention (7,787 rows).")
    add_paragraph(doc, "Preprocessing engineered 16 additional features, including datetime components (year, month, day of week), content age upon Netflix addition, duration metrics (binned and separated into numeric units), country counts, cast counts, primary genre extractions, target audience binnings (Kids, Teens, Adults), and binary type encodings. The resulting datasets, modular Python scripts, reproducible Jupyter Notebook, visual charts, and documentation have been formatted and published to GitHub under the repository yuva-intern-week1-netflix.")

    # ================= 3. INTRODUCTION =================
    add_heading_1(doc, "2. Introduction")
    add_paragraph(doc, "Data acquisition, cleaning, and preprocessing represent the fundamental backbone of modern data science and data engineering workflows. In real-world enterprise environments, raw data is seldom clean, structured, or immediately suitable for machine learning algorithms or executive reporting. Unaddressed missing values, inconsistent formatting, misaligned values, and lack of engineered domain features can lead to biased statistical inferences, degraded model performance, and flawed decision-making.")
    add_paragraph(doc, "The primary objectives of this Week 1 Yuva Internship project are to:")
    add_paragraph(doc, "• Formulate a clean, industry-standard project structure for reproducible data science.\n"
                      "• Acquire authentic raw data from reliable public repositories and preserve raw files without modification.\n"
                      "• Conduct rigorous Exploratory Data Analysis (EDA) to understand feature distributions, patterns, and anomalies.\n"
                      "• Perform comprehensive Data Quality Assessment to identify missingness, duplicates, outliers, and formatting flaws.\n"
                      "• Design and implement an automated Data Cleaning and Preprocessing Pipeline in Python.\n"
                      "• Generate high-resolution, informative visual representations using Matplotlib and Seaborn.\n"
                      "• Publish the entire verified pipeline, code modules, notebook, and report to GitHub.")

    # ================= 4. DATASET DESCRIPTION =================
    add_heading_1(doc, "3. Dataset Description")
    add_paragraph(doc, "The dataset analyzed in this project is the widely acclaimed 'Netflix Movies and TV Shows' dataset created by Shivam Bansal, sourced originally from Flixable and Kaggle, and hosted publicly via TidyTuesday. The raw dataset consists of 7,787 rows and 12 columns, capturing comprehensive catalog metadata for Netflix titles up to early 2021.")

    add_heading_2(doc, "Data Dictionary")
    table_dict = doc.add_table(rows=13, cols=4)
    dict_headers = ["Column Name", "Data Type", "Missing Count (%)", "Description"]
    for i, h in enumerate(dict_headers):
        table_dict.cell(0, i).paragraphs[0].text = h

    dict_data = [
        ("show_id", "String / Object", "0 (0.0%)", "Unique identifier for each movie or TV show entry"),
        ("type", "String / Object", "0 (0.0%)", "Category of content: 'Movie' or 'TV Show'"),
        ("title", "String / Object", "0 (0.0%)", "Official title of the movie or TV show"),
        ("director", "String / Object", "2,389 (30.68%)", "Director(s) of the content"),
        ("cast", "String / Object", "718 (9.22%)", "Main actors and cast members involved"),
        ("country", "String / Object", "507 (6.51%)", "Country or countries where content was produced"),
        ("date_added", "String / Object", "10 (0.13%)", "Date content was added to Netflix catalog"),
        ("release_year", "Integer (int64)", "0 (0.0%)", "Original release year of the content"),
        ("rating", "String / Object", "7 (0.09%)", "Age / audience rating code (e.g., TV-MA, PG-13)"),
        ("duration", "String / Object", "0 (0.0%)", "Duration in minutes (Movies) or Seasons (TV Shows)"),
        ("listed_in", "String / Object", "0 (0.0%)", "Genres and categories separated by commas"),
        ("description", "String / Object", "0 (0.0%)", "Brief textual summary of the plot / premise")
    ]

    for row_idx, data_tuple in enumerate(dict_data, start=1):
        for col_idx, text in enumerate(data_tuple):
            table_dict.cell(row_idx, col_idx).paragraphs[0].text = text
            
    style_table(table_dict)

    # ================= 5. TOOLS AND TECHNOLOGIES =================
    add_heading_1(doc, "4. Tools and Technologies")
    add_paragraph(doc, "The project relies on the standard Python Data Science ecosystem, leveraging modern open-source libraries for computation, data manipulation, visualization, documentation, and version control:")
    add_paragraph(doc, "• Python (v3.13.7): Primary programming language for pipeline scripting and computational analysis.\n"
                      "• Pandas (v2.2.3): High-performance data structure manipulation, missing value handling, and feature extraction.\n"
                      "• NumPy (v2.1.0): Multi-dimensional array operations, vectorized mathematical calculations, and missing value indicators.\n"
                      "• Matplotlib (v3.9.2) & Seaborn (v0.13.2): Statistical visual plotting, customized aesthetic themes, and figure exports.\n"
                      "• Jupyter Notebook & nbformat: Interactive analysis environment and automated notebook generation.\n"
                      "• Python-Docx: Programmatic creation of structured DOCX reports with styled tables and embedded figures.\n"
                      "• Git & GitHub: Version control system and public cloud code repository management.")

    # ================= 6. INITIAL DATA EXPLORATION =================
    add_heading_1(doc, "5. Initial Data Exploration")
    add_paragraph(doc, "Exploratory Data Analysis was executed to establish baseline characteristics of the catalog. The 7,787 content items comprise 5,377 Movies (69.05%) and 2,410 TV Shows (30.95%), indicating that Netflix's legacy catalog is heavily weighted toward feature films.")

    add_image_with_caption(doc, os.path.join(VIS_DIR, '01_content_type_distribution.png'), "Distribution of Content Type (Movies vs TV Shows)")

    add_paragraph(doc, "Temporal analysis of release years shows that while content spans from 1925 to 2021, over 85% of titles were released after 2010, reflecting Netflix's exponential licensing expansion and original programming shift over the last decade.")

    add_image_with_caption(doc, os.path.join(VIS_DIR, '02_release_year_distribution.png'), "Distribution of Catalog Content by Release Year")

    add_paragraph(doc, "Geographic production analysis reveals that the United States is the primary content contributor with 2,877 titles, followed by India (990 titles), the United Kingdom (576 titles), Canada (259 titles), and Japan (235 titles).")

    add_image_with_caption(doc, os.path.join(VIS_DIR, '04_top_countries.png'), "Top 10 Content Producing Countries")

    # ================= 7. DATA QUALITY ASSESSMENT =================
    add_heading_1(doc, "6. Data Quality Assessment")
    add_paragraph(doc, "A comprehensive data audit was conducted across all 12 columns to uncover quality deficiencies. Key findings include:")

    add_image_with_caption(doc, os.path.join(VIS_DIR, '03_missing_values_original.png'), "Missing Value Breakdown in Original Dataset")

    add_paragraph(doc, "1. Missing Values: 'director' exhibited severe missingness (30.68%), primarily due to TV Shows having uncredited or episodic multi-director formats. 'cast' had 9.22% missingness, while 'country' lacked entries for 6.51% of titles. 'date_added' (10 entries) and 'rating' (7 entries) showed minor missingness.\n"
                      "2. Duplicate Records: Zero exact row duplicates were detected. Title duplicate checks revealed duplicate names (e.g., 'The Choice', 'Connect'), but cross-examination confirmed these were distinct movies released in different years or countries.\n"
                      "3. String Formatting: Leading and trailing whitespace was present in text attributes, requiring trimming.\n"
                      "4. Misaligned Attributes: No misaligned duration strings were found in the rating column in this dataset version.")

    # ================= 8. DATA CLEANING METHODOLOGY =================
    add_heading_1(doc, "7. Data Cleaning Methodology")
    add_paragraph(doc, "Data cleaning prioritized data preservation over row deletion. Deleting 30% of rows due to missing director names would severely bias catalog insights. Consequently, explicit imputation strategies were adopted:")

    table_clean = doc.add_table(rows=7, cols=4)
    clean_headers = ["Attribute", "Missing Count", "Cleaning Strategy", "Rationale"]
    for i, h in enumerate(clean_headers):
        table_clean.cell(0, i).paragraphs[0].text = h

    clean_data_rows = [
        ("director", "2,389", "Impute with 'Unknown Director'", "Preserves row integrity while explicitly flagging missing credit"),
        ("cast", "718", "Impute with 'Unknown Cast'", "Maintains record count for non-cast documentaries & animation"),
        ("country", "507", "Impute with 'Unknown Country'", "Prevents loss of international titles with missing metadata"),
        ("rating", "7", "Impute with Mode ('TV-MA')", "Uses statistical mode to fill negligible missing rating codes"),
        ("date_added", "10", "Impute with 'release_year-01-01'", "Provides plausible default date aligned with release year"),
        ("All text fields", "7,787", "Strip leading/trailing spaces", "Standardizes string formatting across all text columns")
    ]

    for row_idx, data_tuple in enumerate(clean_data_rows, start=1):
        for col_idx, text in enumerate(data_tuple):
            table_clean.cell(row_idx, col_idx).paragraphs[0].text = text
            
    style_table(table_clean)

    add_image_with_caption(doc, os.path.join(VIS_DIR, '06_missing_values_before_after.png'), "Missing Values Comparison: Before vs After Data Cleaning")

    # ================= 9. DATA PREPROCESSING METHODOLOGY =================
    add_heading_1(doc, "8. Data Preprocessing Methodology")
    add_paragraph(doc, "Preprocessing transformed the cleaned dataset into an enriched, feature-complete format (28 columns total). Key transformations include:")

    add_paragraph(doc, "1. Datetime Parsing & Temporal Features: 'date_added' was parsed into datetime objects to extract 'year_added', 'month_added', 'month_name_added', 'day_added', and 'day_of_week_added'.\n"
                      "2. Content Age Feature: Created 'content_age' = 'year_added' - 'release_year', quantifying catalog latency.\n"
                      "3. Duration Disaggregation: Split duration into numeric 'duration_int' (minutes for movies, seasons for TV shows) and categorical 'duration_unit'.\n"
                      "4. Target Audience Binning: Categorized ratings into 'Kids' (G, TV-Y, TV-G), 'Teens' (PG, PG-13, TV-14), 'Adults' (R, NC-17, TV-MA), and 'Unrated'.\n"
                      "5. Multi-Valued Attributes: Extracted 'primary_genre', 'genre_count', 'primary_country', 'country_count', and 'cast_count'.\n"
                      "6. Binary Encoding: Added 'is_movie' (1 for Movie, 0 for TV Show).")

    add_image_with_caption(doc, os.path.join(VIS_DIR, '07_addition_month_seasonality.png'), "Catalog Addition Seasonality by Month")
    add_image_with_caption(doc, os.path.join(VIS_DIR, '09_top_primary_genres.png'), "Top 12 Primary Genres on Netflix")
    add_image_with_caption(doc, os.path.join(VIS_DIR, '10_target_audience_distribution.png'), "Target Audience Binning Classification")
    add_image_with_caption(doc, os.path.join(VIS_DIR, '11_movie_duration_distribution.png'), "Distribution of Movie Durations (Minutes)")

    # ================= 10. RESULTS AND EVALUATION =================
    add_heading_1(doc, "9. Results and Evaluation")
    add_paragraph(doc, "The execution of the cleaning and preprocessing pipeline produced measurable improvements in data quality, retention, and analytical readiness:")

    table_eval = doc.add_table(rows=6, cols=3)
    eval_headers = ["Metric / Dimension", "Original Dataset", "Cleaned & Processed Dataset"]
    for i, h in enumerate(eval_headers):
        table_eval.cell(0, i).paragraphs[0].text = h

    eval_rows = [
        ("Total Rows Retained", "7,787", "7,787 (100% Retention)"),
        ("Total Columns", "12", "28 (16 Features Engineered)"),
        ("Total Missing Cells", "3,631", "0 (100% Imputed)"),
        ("Data Format Standardized", "Partial (Strings containing spaces)", "Full (Trimmed & Parsed Datetime)"),
        ("Analysis Readiness", "Low (Requires parsing & imputation)", "High (Ready for ML & Visualization)")
    ]

    for row_idx, data_tuple in enumerate(eval_rows, start=1):
        for col_idx, text in enumerate(data_tuple):
            table_eval.cell(row_idx, col_idx).paragraphs[0].text = text
            
    style_table(table_eval)

    # ================= 11. CHALLENGES AND SOLUTIONS =================
    add_heading_1(doc, "10. Challenges and Solutions")
    add_paragraph(doc, "During project implementation, several technical and data challenges were encountered and successfully resolved:")
    add_paragraph(doc, "• Challenge 1: High missingness in director and cast fields.\n"
                      "  Solution: Implemented explicit domain string flags ('Unknown Director', 'Unknown Cast') to avoid row loss while preserving missingness signals.\n"
                      "• Challenge 2: Heterogeneous duration strings ('90 min' vs '1 Season').\n"
                      "  Solution: Wrote custom regex pattern matching to disaggregate duration into numeric quantities and unit flags.\n"
                      "• Challenge 3: Inconsistent date strings across missing and present entries.\n"
                      "  Solution: Combined release year fallbacks with string-to-datetime parsers to eliminate NaT values.")

    # ================= 12. LIMITATIONS AND FUTURE IMPROVEMENTS =================
    add_heading_1(doc, "11. Limitations and Future Improvements")
    add_paragraph(doc, "While the pipeline achieves complete missing value resolution and rich feature creation, certain limitations exist:")
    add_paragraph(doc, "1. Country Imputation Limitations: Imputing 'Unknown Country' preserves rows but does not recover true origin. Future iterations could scrape IMDb metadata to fill country credits.\n"
                      "2. Text NLP Features: Currently, the 'description' field remains unparsed. Advanced NLP (TF-IDF, Sentiment Analysis, Word Embeddings) can be applied in future weeks for content recommendation models.")

    # ================= 13. CONCLUSION =================
    add_heading_1(doc, "12. Conclusion")
    add_paragraph(doc, "The Week 1 Yuva Internship project successfully achieved all milestone objectives. By transforming raw Netflix metadata into an analysis-ready dataset of 7,787 rows and 28 features, establishing modular Python packages, building an interactive Jupyter Notebook, generating publication-grade visual charts, and deploying to GitHub, this project demonstrates high standards of data engineering, analytical rigor, and technical documentation.")

    # ================= 14. REFERENCES =================
    add_heading_1(doc, "13. References")
    add_paragraph(doc, "1. Bansal, Shivam. (2021). Netflix Movies and TV Shows Dataset. Kaggle. https://www.kaggle.com/datasets/shivamb/netflix-shows\n"
                      "2. R for Data Science TidyTuesday Repository. (2021). Netflix Dataset Release. https://github.com/rfordatascience/tidytuesday\n"
                      "3. Pandas Development Team. (2024). Pandas API Reference and Data Cleaning Guidelines. https://pandas.pydata.org/\n"
                      "4. Matplotlib & Seaborn Documentation. Statistical Data Visualization in Python. https://seaborn.pydata.org/")

    # ================= 15. APPENDIX =================
    add_heading_1(doc, "14. Appendix")
    add_paragraph(doc, "Key Code Snippet: Missing Value Imputation and Text Standardization Routine")

    p_code = doc.add_paragraph()
    p_code.paragraph_format.space_before = Pt(4)
    p_code.paragraph_format.space_after = Pt(10)
    run_code = p_code.add_run(
        "def clean_netflix_data(df):\n"
        "    # Trim spaces across string fields\n"
        "    for col in df.select_dtypes(include=['object']).columns:\n"
        "        df[col] = df[col].astype(str).str.strip().replace({'nan': np.nan, '': np.nan})\n\n"
        "    # Impute missing categoricals\n"
        "    df['director'] = df['director'].fillna('Unknown Director')\n"
        "    df['cast'] = df['cast'].fillna('Unknown Cast')\n"
        "    df['country'] = df['country'].fillna('Unknown Country')\n"
        "    df['rating'] = df['rating'].fillna(df['rating'].mode()[0])\n"
        "    df['date_added'] = df['date_added'].fillna(df['release_year'].astype(str) + '-01-01')\n"
        "    return df"
    )
    run_code.font.name = 'Consolas'
    run_code.font.size = Pt(9)
    run_code.font.color.rgb = RGBColor(30, 41, 59)

    doc.save(OUTPUT_DOCX)
    print(f"Report successfully saved to: {OUTPUT_DOCX}")

if __name__ == '__main__':
    generate_report()
