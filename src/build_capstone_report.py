from pathlib import Path
import json
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
M=json.loads((ROOT/'dataset/processed/week6_model_results.json').read_text(encoding='utf-8'))
Q=json.loads((ROOT/'dataset/processed/week6_data_quality.json').read_text(encoding='utf-8'))
C=json.loads((ROOT/'dataset/processed/week6_model_comparison.csv').read_text(encoding='utf-8').replace('\r','').split('\n')[0]) if False else None
import pandas as pd
comp=pd.read_csv(ROOT/'dataset/processed/week6_model_comparison.csv')
FIG=ROOT/'visualizations'/'week6'; FIG.mkdir(parents=True,exist_ok=True)
# Workflow graphic
fig,ax=plt.subplots(figsize=(11,2.1)); ax.axis('off')
stages=['Source snapshot','Quality checks','EDA + target','Train-only pipeline','CV + test','Evidence-based limits']
for i,s in enumerate(stages):
 x=i*1.8+.1; ax.add_patch(FancyBboxPatch((x,.55),1.48,.72,boxstyle='round,pad=.06',fc=['#e7f0fa','#e8f5e9','#fff3e0','#e8eaf6','#e0f7fa','#fce4ec'][i],ec='#345'))
 ax.text(x+.74,.91,s,ha='center',va='center',fontsize=8,wrap=True)
 if i<len(stages)-1: ax.annotate('',xy=(x+1.78,.91),xytext=(x+1.50,.91),arrowprops={'arrowstyle':'->','color':'#345'})
ax.set_xlim(0,10.9); ax.set_ylim(.35,1.5); fig.tight_layout(); fig.savefig(FIG/'workflow.png',dpi=180,bbox_inches='tight'); plt.close(fig)

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.68); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9.5); styles['Normal'].paragraph_format.space_after=Pt(5)
for nm,size,col in [('Title',28,'17324D'),('Heading 1',17,'17324D'),('Heading 2',12,'277DA1'),('Heading 3',10,'334455')]:
 st=styles[nm]; st.font.name='Aptos Display'; st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(col)
 st.font.bold=True
# footer page number field
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Yuva Internship | Week 6 Capstone  •  Page ')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

def p(text='',style=None,boldlead=None):
 para=doc.add_paragraph(style=style)
 if boldlead and text.startswith(boldlead):
  para.add_run(boldlead).bold=True; para.add_run(text[len(boldlead):])
 else: para.add_run(text)
 return para

def heading(text,level=1): doc.add_heading(text,level=level)
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers)); t.style='Light Shading Accent 1'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for cell,text in zip(t.rows[0].cells,headers): cell.text=str(text); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
 for row in rows:
  cells=t.add_row().cells
  for c,v in zip(cells,row): c.text=str(v); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
 return t

def figcap(name,caption,width=6.3):
 path=FIG/name
 if path.exists():
  doc.add_picture(str(path),width=Inches(width)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
  cp=doc.add_paragraph(caption); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER; cp.runs[0].italic=True; cp.runs[0].font.size=Pt(8)

# cover
for _ in range(2): doc.add_paragraph('')
x=doc.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=x.add_run('YUVA INTERNSHIP'); r.bold=True; r.font.size=Pt(17); r.font.color.rgb=RGBColor(39,125,161)
x=doc.add_paragraph(style='Title'); x.alignment=WD_ALIGN_PARAGRAPH.CENTER; x.add_run('Final Capstone Project')
x=doc.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=x.add_run('End-to-End Data Science Capstone Using Python'); r.bold=True; r.font.size=Pt(16)
x=doc.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER; x.add_run('\nNetflix Catalogue Timing: An Exploratory and Predictive Analysis\n\n')
for label,val in [('Student','Vikhyath Bharadwaj K S'),('Dataset','Netflix Movies and TV Shows (Kaggle snapshot)'),('Submission date','3 October 2026')]:
 x=doc.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER; x.add_run(label+': ').bold=True; x.add_run(val)
doc.add_page_break()
heading('Contents')
for s in ['Abstract','1. Introduction','2. Problem and objectives','3. Dataset and data quality','4. Methodology','5. Cleaning and preprocessing','6. Exploratory analysis','7. Feature engineering','8. Modeling and validation','9. Evaluation and results','10. Findings and recommendations','11. Challenges and limitations','12. Future scope','13. Conclusion','References','Appendix']:
 p(s)
doc.add_page_break()
heading('Abstract')
p(f"This capstone studies a catalogue metadata proxy for the gap between a title’s release year and its recorded addition year in a Netflix catalogue snapshot. The original file contains {M['raw_rows']:,} titles and 12 columns. After requiring a measurable nonnegative target, {M['analysis_rows']:,} rows remained; 98 had an unparseable or missing addition date and 12 had a negative gap. Data exploration showed a strongly right-skewed proxy (median 1 year; 90th percentile 13 years). A scikit-learn pipeline compared a median baseline, Ridge regression, Random Forest, and histogram gradient boosting using five-fold cross-validation on an 80% training split. Random Forest was selected by mean CV MAE and achieved held-out MAE {M['test_mae']:.2f} years, RMSE {M['test_rmse']:.2f} years, and R-squared {M['test_r2']:.3f} on a {M['test_rows']:,}-row holdout. Metadata offers some predictive signal over the median baseline, but error remains material. The results are associational and do not measure licensing duration, viewer demand, or causal mechanisms.")
heading('1. Introduction')
p('A catalogue contains a mixture of titles released in different years and recorded as added at different times. Describing that difference can help characterize the catalogue’s historical composition. The available dataset supports a limited metadata analysis; it does not contain licensing, customer, or viewing outcomes.')
p('The capstone deliberately poses one question distinct from the earlier type-classification and clustering tasks: how well can metadata estimate an approximate release-to-addition year gap? The work combines data quality checks, exploratory plots, a training-only preprocessing pipeline, candidate regression models, held-out evaluation, and cautious interpretation.')
heading('2. Problem and objectives')
p('Problem statement: estimate the whole-year difference between `date_added` year and `release_year` from other title metadata, and describe how that proxy varies in the supplied snapshot.')
for s in ['Audit relevant field completeness and identify impossible proxy values.','Describe the delay distribution and selected catalogue patterns.','Compare candidate regressors against a median baseline using cross-validation and a reserved test set.','Explain what the model can and cannot support.']: p(s,style='List Bullet')
p('Research questions: How skewed is the proxy? How does it vary among common listed genres and by addition year? Do metadata predictors lower held-out MAE compared with the median baseline?')
heading('3. Dataset and data quality')
p('The local CSV is attributed to Shivam Bansal and listed at Kaggle. It has 7,787 rows and 12 fields: show identifier, type, title, director, cast, country, date added, release year, rating, duration, listed genres, and description. Data are catalogue metadata. The Kaggle page is the source of the snapshot; the publisher does not document a row-level sampling method or a complete snapshot timestamp in this project.')
p('The untouched source has no exact duplicate rows and no missing release years. The raw `date_added` field has 10 null cells, with additional values that do not parse as dates, leaving 98 unusable addition-date fields for this calculation. Other raw missingness includes director (2,389), cast (718), country (507), and rating (7). Missing fields are handled inside the estimator where selected as predictors.')
table(['Quality check','Result / action'],[['Raw dimensions',f"{M['raw_rows']:,} × 12"],['Exact duplicate rows',M['exact_duplicate_rows']],['Missing/unparseable date_added',M['missing_date_added_rows']],['Missing release_year',M['missing_release_year_rows']],['Negative proxy values',f"{M['negative_delay_rows_excluded']} excluded from target analysis"],['Modeling population',f"{M['analysis_rows']:,} valid nonnegative values"]])
heading('4. Methodology',1)
p('The workflow retains the original source, parses fields, builds a transparent target proxy, explores the valid analysis subset, then fits transformations and regressors on training data only. Five-fold shuffled K-fold cross-validation compares models in the training split; the reserved 20% test split is used for final evaluation.')
figcap('workflow.png','Figure 1. Capstone workflow from local snapshot through evaluation and interpretation.')
heading('5. Cleaning and preprocessing')
p('The pipeline creates a parsed `date_added`, numeric `release_year`, and numeric `duration_value` from the leading integer in duration. It derives primary genre and country from the first comma-separated value, plus genre and country counts. Missing multi-valued fields map to an “Unknown” primary category and zero count. This preserves a simple, reproducible representation while losing information in the full multi-valued fields.')
p('The target is `year(date_added) - release_year`. It is measured in whole calendar years, not exact elapsed time. Ninety-eight rows could not be assigned the proxy because addition dates were missing or failed parsing; 12 negative differences are inconsistent with this proxy’s intended interpretation and are excluded only from target analysis. No nonnegative high-delay observation is removed as an outlier. The raw file is left unchanged.')
p('The operation is intentionally not described as correcting source data: the invalid rows remain in the original CSV and are transparently excluded from this particular model. The 93-year maximum remains in the analysis; MAE is emphasized because squared error is more sensitive to the long tail.')
heading('6. Exploratory analysis')
p(f"Among {M['analysis_rows']:,} valid rows, the mean proxy is {Q['delay_describe']['mean']:.2f} years, median 1 year, 75th percentile 5 years, 90th percentile 13 years, and maximum 93 years. The distribution is strongly right-skewed: at least a quarter of titles have a zero calendar-year difference.")
figcap('delay_distribution.png','Figure 2. Approximate calendar-year difference for titles with a valid, nonnegative proxy.')
figcap('delay_by_genre.png','Figure 3. Distribution by common primary listed genre; extreme points are hidden only to keep the plot readable.')
figcap('additions_by_year.png','Figure 4. Snapshot counts by recorded addition year; counts reflect the dataset snapshot and coverage.')
p('The grouped genre chart compares first-listed genre labels and should not be read as a causal genre effect. The addition-year chart describes counts represented in this file; incomplete historical capture and snapshot design limit interpretation as a catalogue growth series.')
heading('7. Feature engineering')
p('Predictors are type, rating, primary genre, primary country, genre count, country count, and duration value. `year_added` and `release_year` are explicitly excluded because together they calculate the target. Identifiers, title, cast, director, and free-text description are also excluded. Categorical missing values use most-frequent imputation, rare categories are grouped by the encoder, and one-hot encoding handles unknown categories. Numeric fields are median-imputed and standardized. Each transform is fitted within the pipeline during each training fold.')
heading('8. Modeling and validation')
p('The baseline predicts the training-fold median. Candidate models are Ridge regression (alpha 10), Random Forest (80 trees, minimum leaf 3, max_features 0.8, seed 42), and histogram gradient boosting (160 iterations, learning rate 0.08, L2 regularization 1.0). The Random Forest was selected by lower mean cross-validated MAE. These fixed settings provide a bounded comparison rather than exhaustive tuning.')
p('The data were split randomly 80/20 with seed 42. Five-fold shuffled K-fold cross-validation was performed on training data. This procedure avoids fitting preprocessing on held-out folds, although random splitting does not test future-year generalization.')
heading('9. Evaluation and results')
p('MAE gives the average absolute error in years; RMSE gives more weight to large errors; R-squared compares model error with the mean-target baseline. The selected Random Forest improves held-out MAE by about 0.45 years relative to the median baseline, but its typical error remains several years.')
rows=[]
for _,r in comp.iterrows(): rows.append([r['model'],f"{r['cv_mae_mean']:.2f} ± {r['cv_mae_sd']:.2f}",f"{r['test_mae']:.2f}",f"{r['test_rmse']:.2f}",f"{r['test_r2']:.3f}"])
table(['Model','CV MAE (years)','Test MAE','Test RMSE','Test R²'],rows)
figcap('model_comparison.png','Figure 5. Candidate model comparison by held-out MAE; smaller is better.')
figcap('observed_vs_estimated.png','Figure 6. Observed versus predicted proxy; axes show the central 99th-percentile range.')
p(f"Selected Random Forest test metrics: MAE {M['test_mae']:.2f} years, RMSE {M['test_rmse']:.2f} years, and R-squared {M['test_r2']:.3f}. These results do not justify individual operational decisions. The target’s long tail and coarse calendar-year construction create substantial uncertainty.")
heading('10. Findings and recommendations')
for s in ['Observed: the proxy is highly right-skewed; the median is one year while the 90th percentile is 13 years.','Observed: selected metadata improves average absolute error compared with the median-only baseline on this random split.','Model-based: the selected forest has held-out MAE near 3.76 years and R-squared near 0.36; substantial variation remains unexplained.','Recommendation: use the plots and estimates only for exploratory catalogue description. Do not infer terms, popularity, or audience response.']: p(s,style='List Bullet')
p('Any operational application would require a representative, current panel and independently measured outcomes relevant to the decision. Associations in this single snapshot do not identify causes.')
heading('11. Challenges and limitations')
for s in ['Missing or unparseable addition dates prevent target calculation for 98 titles; 12 rows had negative calendar-year gaps.','Year subtraction is coarse: exact dates and within-year timing are ignored.','The source snapshot’s collection process and coverage are not fully documented here.','The first listed genre/country summarize multi-valued columns in an order-dependent way.','Random holdout validation may overstate performance for later catalogue periods.','The target is not licensing duration and predictor associations are not causal.','The long tail makes RMSE sensitive to a small number of large errors; MAE is reported as the primary measure.']: p(s,style='List Bullet')
heading('12. Future scope')
p('Future work could use repeated dated snapshots to test temporal generalization, a more complete multi-label representation, subgroup error analysis, prediction intervals, and external variables only where licensing and usage rights permit. A time-based validation split would be necessary before claiming prospective performance.')
heading('13. Conclusion')
p(f"The project completes an end-to-end, reproducible analysis of catalogue timing metadata. The Random Forest outperformed a median baseline on the held-out split, with a test MAE of {M['test_mae']:.2f} years. Since the outcome is only a calendar-year proxy and error remains material, the strongest conclusion is descriptive: catalogue metadata contains limited signal about this gap, while the source does not support causal or licensing claims.")
heading('References')
for s in ['Bansal, S. Netflix Movies and TV Shows. Kaggle. https://www.kaggle.com/datasets/shivamb/netflix-shows (accessed 3 October 2026).','scikit-learn User Guide. https://scikit-learn.org/stable/user_guide.html','pandas documentation. https://pandas.pydata.org/docs/','Matplotlib documentation. https://matplotlib.org/stable/','Seaborn documentation. https://seaborn.pydata.org/']: p(s,style='List Number')
heading('Appendix')
heading('Selected implementation details',2)
p('Target construction: `catalogue_delay_years = date_added_parsed.dt.year - release_year`. Model predictors: type, rating, primary genre/country, genre count, country count, and duration value. Preprocessing is a scikit-learn ColumnTransformer within a Pipeline so transformations are fitted only to training data.')
p('Source code: `src/capstone_pipeline.py`. Notebook: `notebooks/capstone_data_science_project.ipynb`. Run-specific comparison and diagnostics are saved in `dataset/processed/week6_model_comparison.csv`, `week6_model_results.json`, and `week6_data_quality.json`.')
path=ROOT/'reports'/'Yuva_Internship_Final_Capstone_Report.docx'; doc.save(path); print(path)
