import json
import os
from io import BytesIO
from datetime import date

import docx
from PIL import Image
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

PROJECT_DIR = os.path.dirname(os.path.dirname(__file__))
REPORT_PATH = os.path.join(PROJECT_DIR, "reports", "Week3_Clustering_Analysis_Report.docx")
VIS_DIR = os.path.join(PROJECT_DIR, "visualizations", "week3")
METRICS_PATH = os.path.join(PROJECT_DIR, "dataset", "processed", "week3_clustering_metrics.json")

COLOR_PRIMARY = RGBColor(229, 9, 20)
COLOR_SECONDARY = RGBColor(34, 31, 31)
COLOR_NAVY = RGBColor(26, 82, 118)
COLOR_TEXT = RGBColor(51, 51, 51)


def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcPr.append(parse_xml(
        f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>'
    ))


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY


def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = COLOR_NAVY


def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = COLOR_SECONDARY


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = COLOR_TEXT
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = COLOR_TEXT


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = COLOR_TEXT
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = COLOR_TEXT


def add_callout(doc, title, text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F2F4F4")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    tcPr = cell._element.get_or_add_tcPr()
    tcPr.append(parse_xml(
        f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="36" w:space="0" w:color="E50914"/>'
        f"<w:top w:val='none'/><w:right w:val='none'/><w:bottom w:val='none'/></w:tcBorders>"
    ))
    p = cell.paragraphs[0]
    rt = p.add_run(f"{title}\n")
    rt.font.name = "Calibri"
    rt.font.size = Pt(11)
    rt.font.bold = True
    rt.font.color.rgb = COLOR_PRIMARY
    rb = p.add_run(text)
    rb.font.name = "Calibri"
    rb.font.size = Pt(10.5)
    rb.font.color.rgb = COLOR_TEXT
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_figure(doc, filename, fig_num, title, caption, width=6.0):
    path = os.path.join(VIS_DIR, filename)
    if not os.path.exists(path):
        add_body(doc, f"[Missing figure: {filename}]")
        return
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.keep_with_next = True
    # Embed compact, print-readable JPEGs so the submission DOCX stays under
    # the portal's 2 MB upload limit. Keep the full-resolution PNGs in the repo.
    with Image.open(path) as source:
        image = source.convert("RGBA")
        image.thumbnail((1450, 1100), Image.Resampling.LANCZOS)
        background = Image.new("RGB", image.size, "white")
        background.paste(image, mask=image.getchannel("A"))
        compact = BytesIO()
        background.save(compact, format="JPEG", quality=78, optimize=True, progressive=True)
        compact.seek(0)
    p_img.add_run().add_picture(compact, width=Inches(width))
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_after = Pt(10)
    r1 = p_cap.add_run(f"Figure {fig_num}: ")
    r1.font.name = "Calibri"
    r1.font.size = Pt(9.5)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_PRIMARY
    r2 = p_cap.add_run(f"{title} — ")
    r2.font.name = "Calibri"
    r2.font.size = Pt(9.5)
    r2.font.bold = True
    r2.font.color.rgb = COLOR_SECONDARY
    r3 = p_cap.add_run(caption)
    r3.font.name = "Calibri"
    r3.font.size = Pt(9.5)
    r3.font.italic = True
    r3.font.color.rgb = COLOR_TEXT


def create_styled_table(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, title in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        cell.text = title
        set_cell_background(cell, "1A5276")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(9)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for r_idx, row in enumerate(data):
        bg = "F9FAFC" if r_idx % 2 else "FFFFFF"
        for c_idx, val in enumerate(row):
            cell = tbl.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(9)
                run.font.color.rgb = COLOR_TEXT
    if col_widths:
        for row in tbl.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def fmt(x, nd=3):
    if isinstance(x, float):
        return f"{x:.{nd}f}"
    return str(x)


def pct_join(d, n=3):
    items = sorted(d.items(), key=lambda kv: -kv[1])[:n]
    return "; ".join(f"{k} {v:.1f}%" for k, v in items)


def generate_report():
    with open(METRICS_PATH, encoding="utf-8") as f:
        m = json.load(f)
    movies = m["movies"]
    tv = m["tv_shows"]
    cmp_ = m["comparison"]
    mp = {p["cluster"]: p for p in movies["profiles"]}
    tp = {p["cluster"]: p for p in tv["profiles"]}

    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    r = p.add_run("YUVA INTERNSHIP PROGRAM — WEEK 3 DELIVERABLE")
    r.font.name = "Calibri"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = COLOR_PRIMARY

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("Unsupervised Learning and Clustering Analysis")
    r.font.name = "Calibri"
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    p.style = doc.styles["Title"]
    r.font.color.rgb = RGBColor(0, 0, 0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run("Netflix Movies and TV Shows — K-Means and Hierarchical Clustering of Catalogue Metadata")
    r.font.name = "Calibri"
    r.font.size = Pt(13)
    r.font.italic = True
    r.font.color.rgb = COLOR_TEXT

    create_styled_table(doc, ["Metadata Field", "Details"], [
        ["Student Name", "Vikhyath Bharadwaj K S"],
        ["Project Title", "Netflix Movies and TV Shows — Unsupervised Learning and Clustering Analysis"],
        ["Dataset", "Netflix Movies and TV Shows by Shivam Bansal (Kaggle)"],
        ["Records analysed", "7,787 titles (5,377 Movies; 2,410 TV Shows)"],
        ["Primary models", "K-Means, K=4 (Movies) and K=4 (TV Shows); random_state=42"],
        ["Environment", "Python | Pandas | NumPy | Matplotlib | Seaborn | Scikit-learn | SciPy"],
        ["Submission Date", date.today().strftime("%B %d, %Y")],
    ], col_widths=[2.4, 4.1])

    doc.add_page_break()
    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run("Yuva Internship Week 3 | Vikhyath Bharadwaj K S |")
    fr.font.name = "Calibri"
    fr.font.size = Pt(9)
    fr.font.color.rgb = COLOR_TEXT
    add_page_number(fp)

    add_heading_1(doc, "Table of Contents")
    for item in [
        "1. Abstract", "2. Introduction", "3. Dataset Description", "4. Tools and Technologies",
        "5. Methodology", "6. Data Preprocessing", "7. K-Means Clustering", "8. Hierarchical Clustering",
        "9. Cluster Visualization", "10. Cluster Analysis and Interpretation", "11. Model Evaluation and Comparison",
        "12. Key Findings", "13. Challenges and Solutions", "14. Limitations and Future Scope",
        "15. Conclusion", "16. References", "17. Appendix", "18. Validation Checklist",
    ]:
        add_body(doc, item)

    add_heading_1(doc, "1. Abstract")
    add_body(
        doc,
        "This report applies unsupervised clustering to the cleaned Netflix Movies and TV Shows catalogue "
        "prepared in Weeks 1–2. The objective is to group titles by available content metadata — release timing, "
        "type-specific length, audience band, frequent genres, and grouped production country — without using "
        "show identifiers or titles as clustering features."
    )
    add_body(
        doc,
        f"Movies and TV Shows were modelled separately because runtime minutes and season counts are not "
        f"commensurate. Numeric fields were scaled with RobustScaler; audience and country groups were one-hot "
        f"encoded; frequent genres were multi-hot encoded from listed_in. K-Means (n_init=10, random_state=42) "
        f"was evaluated for K=2…10 using inertia, silhouette, Davies–Bouldin, Calinski–Harabasz, size balance, "
        f"and interpretability. Silhouette is highest at K=2 for both subsets because a compact older-catalogue "
        f"minority separates from a contemporary majority. K=4 is reported as a practical descriptive resolution, "
        f"not a unique optimum: the elbow is gradual, and the extra clusters expose interpretable format and "
        f"genre/country patterns. Movie K=4 silhouette = {fmt(movies['final']['silhouette'])}; TV K=4 silhouette = "
        f"{fmt(tv['final']['silhouette'])}. Ward hierarchical clustering on stratified samples of 250 titles "
        f"shows moderate agreement with K-Means (movie ARI = {fmt(cmp_['movie_hier_ari'])}; TV ARI = "
        f"{fmt(cmp_['tv_hier_ari'])}). Clusters describe catalogue metadata structure only; they are not evidence "
        f"of Netflix’s recommender or of viewer preferences."
    )

    add_heading_1(doc, "2. Introduction")
    add_body(
        doc,
        "Supervised learning predicts a labelled target. Unsupervised learning has no such target: the algorithm "
        "must recover groups, densities, or components from the feature geometry alone. Clustering is the branch "
        "that assigns each observation to a group so that within-group similarity is high relative to between-group "
        "similarity under a chosen distance."
    )
    add_body(
        doc,
        "Clustering is useful in data analysis when labels are unavailable or expensive, when one needs a compact "
        "taxonomy of records, or when one wants to audit whether engineered features actually separate known "
        "domain concepts. In a streaming catalogue, metadata clustering can surface recency bands, format length "
        "profiles, and genre–geography co-occurrence — always as properties of the table, not of hidden user data."
    )
    add_heading_2(doc, "2.1 Assignment objectives")
    add_bullet(doc, "Prepare a justified, leakage-free feature matrix from the Week 1 processed table.", "1. ")
    add_bullet(doc, "Implement K-Means with a documented search over K and multiple validation indices.", "2. ")
    add_bullet(doc, "Implement agglomerative clustering with a readable dendrogram via sampling if required.", "3. ")
    add_bullet(doc, "Visualise and interpret clusters using original-scale statistics and descriptive labels.", "4. ")
    add_heading_2(doc, "2.2 Relevance of the dataset")
    add_body(
        doc,
        "The Kaggle Netflix titles table is a public catalogue snapshot (author: Shivam Bansal). Week 2 showed "
        "strong type, country, genre, and duration structure. Those findings motivate clustering but also warn "
        "against mixing movie minutes with TV seasons. The dataset contains no watch history."
    )

    add_heading_1(doc, "3. Dataset Description")
    add_body(
        doc,
        "Source: Kaggle, Netflix Movies and TV Shows, Shivam Bansal. Original size 7,787 × 12. Week 1 retained "
        "all rows and engineered temporal, duration, genre, country, and audience fields (processed size 7,787 × 28). "
        "This week uses that processed file."
    )
    create_styled_table(doc, ["Attribute", "Role in clustering", "Notes"], [
        ["show_id, title", "Interpretation only", "Never used as numeric features"],
        ["type", "Split key", "Separate Movie and TV Show models"],
        ["release_year, content_age", "Release year and addition lag", "Robust-scaled; distinct timing dimensions"],
        ["duration_int", "Length within type", "Minutes for movies; seasons for TV"],
        ["genre_count, country_count", "Numeric complexity", "Robust-scaled"],
        ["target_audience", "One-hot", "Kids / Teens / Adults / Unrated"],
        ["primary_country grouped", "One-hot", "US, India, UK, Unknown, Other"],
        ["listed_in (frequent tags)", "Multi-hot", "12 movie genres; 10 TV genres"],
    ], col_widths=[2.1, 1.8, 2.6])
    add_body(
        doc,
        "Limitations carried forward: unknown country/cast/director placeholders; 10 unparsed date_added values "
        "already handled in content_age; multi-valued countries reduced to a primary country plus a count; 42 raw "
        "genre tags compressed to frequent indicators."
    )

    add_heading_1(doc, "4. Tools and Technologies")
    add_bullet(doc, "end-to-end scripting, notebook execution, and report generation.", "Python: ")
    add_bullet(doc, "table loading, subsetting, and cluster characterisation.", "Pandas: ")
    add_bullet(doc, "numeric matrices and summary calculations.", "NumPy: ")
    add_bullet(doc, "evaluation charts and dendrograms.", "Matplotlib / Seaborn: ")
    add_bullet(doc, "KMeans, AgglomerativeClustering, PCA, RobustScaler, and cluster metrics.", "Scikit-learn: ")
    add_bullet(doc, "Ward linkage and dendrogram construction.", "SciPy: ")
    add_bullet(doc, "documented, executable analysis narrative.", "Jupyter Notebook: ")

    add_heading_1(doc, "5. Methodology")
    add_body(doc, "The workflow is linear and reproducible:")
    add_bullet(doc, "Load netflix_processed.csv from Week 1.", "1. Ingest — ")
    add_bullet(doc, "Split by type; drop identifiers from the feature matrix.", "2. Scope — ")
    add_bullet(doc, "Encode audience/country/genres; Robust-scale numeric columns.", "3. Transform — ")
    add_bullet(doc, "Fit K-Means for K=2…10; compute inertia, silhouette, DB, CH, size ratios.", "4. Select K — ")
    add_bullet(doc, "Fit K=4; attach labels to show_id/title; profile original-scale features.", "5. Fit & profile — ")
    add_bullet(doc, "Ward/average/complete on a 250-title stratified sample; dendrogram.", "6. Hierarchical — ")
    add_bullet(doc, "PCA scatter, centroid heatmap, genre/audience charts; compare algorithms.", "7. Visualise & compare — ")
    add_callout(
        doc, "No combined duration model",
        "A single K-Means run on all 7,787 rows with a shared duration column would treat 90 minutes and 3 seasons "
        "as nearby numbers. That design was rejected. Type is used as a hard split, not as a dummy beside mixed length."
    )

    add_heading_1(doc, "6. Data Preprocessing")
    add_heading_2(doc, "6.1 Feature matrix")
    add_body(
        doc,
        f"Movie matrix: {movies['n']} rows × {movies['n_features']} columns. "
        f"TV matrix: {tv['n']} rows × {tv['n_features']} columns. Binary columns were left as 0/1."
    )
    add_heading_2(doc, "6.2 Missing values")
    add_body(
        doc,
        "Clustering numeric and encoded fields have zero missing values after Week 1. Unknown Country remains an "
        "explicit group so imputed geography is visible rather than silently merged into the United States mode."
    )
    add_heading_2(doc, "6.3 Encoding")
    add_body(
        doc,
        "One-hot encoding is used where categories have no defensible numeric order (country group; Unrated audience). "
        "Multi-hot encoding represents co-listed genres without forcing a single primary_genre. Ordinal encoding of "
        "raw MPAA/TV ratings was avoided because Unrated/NR/UR and the Kids–Teens–Adults collapse are not a single interval scale."
    )
    add_heading_2(doc, "6.4 Scaling")
    add_body(
        doc,
        "RobustScaler subtracts the median and divides by the IQR. It is preferred to StandardScaler here because "
        "content age and movie runtime are skewed and have outliers. Release year and content age capture original "
        "release timing and catalogue-addition lag, respectively. MinMaxScaler would still stretch with "
        "those maxima. Scaling is fit within each type subset only (no leakage from TV into movie geometry)."
    )
    add_heading_2(doc, "6.5 Code excerpt")
    p = doc.add_paragraph()
    run = p.add_run(
        "numeric_cols = ['release_year','duration_int','content_age','genre_count','country_count']\n"
        "scaler = RobustScaler()\n"
        "X_num = scaler.fit_transform(subset[numeric_cols])\n"
        "X = np.hstack([X_num, audience_one_hot, country_one_hot, genre_multi_hot])\n"
        "km = KMeans(n_clusters=4, random_state=42, n_init=10)"
    )
    run.font.name = "Courier New"
    run.font.size = Pt(9)
    run.font.color.rgb = COLOR_NAVY

    add_heading_1(doc, "7. K-Means Clustering")
    add_heading_2(doc, "7.1 Algorithm")
    add_body(
        doc,
        "K-Means maintains K centroids. Each title is assigned to the nearest centroid in Euclidean space; centroids "
        "are then replaced by the mean of their assigned points. The objective is inertia (within-cluster sum of squares). "
        "k-means++ initialisation and n_init=10 reduce empty-cluster and unlucky-seed risk. Limitations include the "
        "implicit spherical cluster assumption, the requirement to choose K, and distortion from mixed binary features."
    )
    add_heading_2(doc, "7.2 Configuration")
    create_styled_table(doc, ["Setting", "Value"], [
        ["Algorithm", "sklearn.cluster.KMeans"],
        ["Candidate K", "2 through 10"],
        ["Reported K (Movies / TV)", f"{movies['selected_k']} / {tv['selected_k']}"],
        ["random_state", str(m["random_state"])],
        ["n_init", str(m["n_init"])],
        ["Distance (implied)", "Euclidean on the scaled mixed matrix"],
        ["Silhouette-best K (both subsets)", "2"],
    ], col_widths=[2.6, 3.9])
    add_heading_2(doc, "7.3 Elbow method")
    add_body(
        doc,
        f"Movie inertia falls from {movies['k_metrics'][0]['inertia']:,.0f} (K=2) to {movies['k_metrics'][2]['inertia']:,.0f} (K=4) to {movies['k_metrics'][3]['inertia']:,.0f} (K=5). "
        f"TV inertia falls from {tv['k_metrics'][0]['inertia']:,.0f} (K=2) to {tv['k_metrics'][2]['inertia']:,.0f} (K=4) to {tv['k_metrics'][3]['inertia']:,.0f} (K=5). "
        "The gains taper gradually; there is no uniquely sharp elbow."
    )
    add_figure(doc, "w3_movie_01_elbow.png", "1", "Movie elbow curve",
               "Inertia versus K for movies. Blue marker: reported K=4; orange would indicate silhouette-best K=2 where plotted.")
    add_figure(doc, "w3_tv_01_elbow.png", "2", "TV Show elbow curve",
               "Inertia versus K for TV Shows with the reported K=4 highlighted.")
    add_heading_2(doc, "7.4 Silhouette and additional indices")
    add_body(
        doc,
        f"Movie silhouette peaks at K=2 ({fmt(movies['k_metrics'][0]['silhouette'])}) and is {fmt(movies['final']['silhouette'])} at K=4. "
        f"TV silhouette peaks at K=2 ({fmt(tv['k_metrics'][0]['silhouette'])}) and is {fmt(tv['final']['silhouette'])} at K=4. "
        f"At K=2, the smaller/larger cluster sizes are movies: {movies['k_metrics'][0]['min_cluster_size']:,}/{movies['k_metrics'][0]['max_cluster_size']:,}; "
        f"TV: {tv['k_metrics'][0]['min_cluster_size']:,}/{tv['k_metrics'][0]['max_cluster_size']:,}. That recency split is real but coarse. "
        f"Davies–Bouldin and Calinski–Harabasz are reported alongside size ratios so that a single index cannot dictate K."
    )
    add_figure(doc, "w3_movie_02_silhouette.png", "3", "Movie silhouette vs K",
               "Mean silhouette coefficient for movie K-Means. High K=2 scores reflect the older-catalogue minority.")
    add_figure(doc, "w3_tv_02_silhouette.png", "4", "TV silhouette vs K",
               "Mean silhouette coefficient for TV K-Means.")
    add_figure(doc, "w3_movie_03_db_ch.png", "5", "Movie Davies–Bouldin and Calinski–Harabasz",
               "Complementary indices for movies. Lower Davies–Bouldin and higher Calinski–Harabasz are preferred.")
    add_figure(doc, "w3_tv_03_db_ch.png", "6", "TV Davies–Bouldin and Calinski–Harabasz",
               "Complementary indices for TV Shows.")
    add_heading_2(doc, "7.5 Reported K-Means results")
    create_styled_table(doc, ["Subset", "K", "Inertia", "Silhouette", "Davies–Bouldin", "Calinski–Harabasz"], [
        ["Movies (full)", "4", fmt(movies["final"]["inertia"], 1), fmt(movies["final"]["silhouette"]),
         fmt(movies["final"]["davies_bouldin"]), fmt(movies["final"]["calinski_harabasz"], 1)],
        ["TV Shows (full)", "4", fmt(tv["final"]["inertia"], 1), fmt(tv["final"]["silhouette"]),
         fmt(tv["final"]["davies_bouldin"]), fmt(tv["final"]["calinski_harabasz"], 1)],
    ], col_widths=[1.5, 0.6, 1.1, 1.1, 1.3, 1.5])
    add_figure(doc, "w3_movie_04_cluster_sizes.png", "7", "Movie cluster sizes",
               "K=4 movie assignment counts and percentages.")
    add_figure(doc, "w3_tv_04_cluster_sizes.png", "8", "TV cluster sizes",
               "K=4 TV assignment counts. The small legacy group is a minority and should be interpreted cautiously.")

    add_heading_1(doc, "8. Hierarchical Clustering")
    add_body(
        doc,
        "Agglomerative clustering merges nearest clusters until one tree remains. Ward linkage minimises the increase "
        "in variance and is paired with Euclidean distance. Average linkage uses mean pairwise distance; complete linkage "
        "uses the farthest pair. A full 5,377-leaf dendrogram is not readable, so each type uses a stratified sample of "
        f"{movies['hierarchical']['sample_size']} titles (stratified by target_audience, seed 42)."
    )
    add_callout(
        doc, "Sampling limitation",
        movies["hierarchical"]["sampling_note"]
    )
    create_styled_table(doc, ["Subset / linkage", "Silhouette (sample)", "Davies–Bouldin", "Calinski–Harabasz"], [
        ["Movies Ward", fmt(movies["hierarchical"]["linkage_metrics"]["ward"]["silhouette"]),
         fmt(movies["hierarchical"]["linkage_metrics"]["ward"]["davies_bouldin"]),
         fmt(movies["hierarchical"]["linkage_metrics"]["ward"]["calinski_harabasz"], 1)],
        ["Movies Average", fmt(movies["hierarchical"]["linkage_metrics"]["average"]["silhouette"]),
         fmt(movies["hierarchical"]["linkage_metrics"]["average"]["davies_bouldin"]),
         fmt(movies["hierarchical"]["linkage_metrics"]["average"]["calinski_harabasz"], 1)],
        ["Movies Complete", fmt(movies["hierarchical"]["linkage_metrics"]["complete"]["silhouette"]),
         fmt(movies["hierarchical"]["linkage_metrics"]["complete"]["davies_bouldin"]),
         fmt(movies["hierarchical"]["linkage_metrics"]["complete"]["calinski_harabasz"], 1)],
        ["TV Ward", fmt(tv["hierarchical"]["linkage_metrics"]["ward"]["silhouette"]),
         fmt(tv["hierarchical"]["linkage_metrics"]["ward"]["davies_bouldin"]),
         fmt(tv["hierarchical"]["linkage_metrics"]["ward"]["calinski_harabasz"], 1)],
        ["TV Average", fmt(tv["hierarchical"]["linkage_metrics"]["average"]["silhouette"]),
         fmt(tv["hierarchical"]["linkage_metrics"]["average"]["davies_bouldin"]),
         fmt(tv["hierarchical"]["linkage_metrics"]["average"]["calinski_harabasz"], 1)],
        ["TV Complete", fmt(tv["hierarchical"]["linkage_metrics"]["complete"]["silhouette"]),
         fmt(tv["hierarchical"]["linkage_metrics"]["complete"]["davies_bouldin"]),
         fmt(tv["hierarchical"]["linkage_metrics"]["complete"]["calinski_harabasz"], 1)],
    ], col_widths=[2.0, 1.7, 1.5, 1.7])
    add_body(
        doc,
        "Average (and, for TV, complete) linkage can outscore Ward on silhouette by isolating tiny compact groups. "
        "Ward is retained for the dendrogram because its clusters are closer in spirit to K-Means variance partitions "
        f"and because movie ARI(K-Means, Ward) = {fmt(cmp_['movie_hier_ari'])} on the sample."
    )
    add_figure(doc, "w3_movie_10_dendrogram.png", "9", "Movie Ward dendrogram",
               "Hierarchical tree for 250 sampled movies. Height is Ward distance, not a catalogue-wide guarantee.")
    add_figure(doc, "w3_tv_10_dendrogram.png", "10", "TV Ward dendrogram",
               "Hierarchical tree for 250 sampled TV titles.")

    add_heading_1(doc, "9. Cluster Visualization")
    add_body(
        doc,
        f"Movie PCA retains {movies['pca']['pc1']*100:.1f}% + {movies['pca']['pc2']*100:.1f}% = "
        f"{movies['pca']['cumulative_2d']*100:.1f}% of scaled-matrix variance in two components. "
        f"TV PCA retains {tv['pca']['pc1']*100:.1f}% + {tv['pca']['pc2']*100:.1f}% = "
        f"{tv['pca']['cumulative_2d']*100:.1f}%. Overlap in the scatter therefore does not prove that original-space "
        "centroids are identical."
    )
    add_figure(doc, "w3_movie_05_pca.png", "11", "Movie PCA scatter",
               "Two-dimensional projection of movie K-Means labels. A 2D view cannot preserve all pairwise distances.")
    add_figure(doc, "w3_tv_05_pca.png", "12", "TV PCA scatter",
               "Two-dimensional projection of TV K-Means labels.")
    add_figure(doc, "w3_movie_06_centroids.png", "13", "Movie scaled centroids",
               "Heatmap of K-Means centres in the scaled movie feature space.")
    add_figure(doc, "w3_tv_06_centroids.png", "14", "TV scaled centroids",
               "Heatmap of K-Means centres in the scaled TV feature space.")
    add_figure(doc, "w3_movie_07_boxplots.png", "15", "Movie numeric distributions",
               "Duration, release year, and content age by movie cluster (original units).")
    add_figure(doc, "w3_tv_07_boxplots.png", "16", "TV numeric distributions",
               "Season count, release year, and content age by TV cluster.")
    add_figure(doc, "w3_movie_08_audience.png", "17", "Movie audience mix",
               "Stacked audience-band percentages by movie cluster.")
    add_figure(doc, "w3_movie_09_genres.png", "18", "Movie genre presence",
               "Percentage of titles in each movie cluster that list selected genres.")
    add_figure(doc, "w3_tv_08_audience.png", "19", "TV audience mix",
               "Stacked audience-band percentages by TV cluster.")
    add_figure(doc, "w3_tv_09_genres.png", "20", "TV genre presence",
               "Percentage of titles in each TV cluster that list selected genres.")
    add_figure(doc, "w3_movie_11_hier_pca.png", "21", "Movie hierarchical sample PCA",
               "Ward labels on the 250-movie sample projected with PCA.")
    add_figure(doc, "w3_compare_01_method_silhouette.png", "22", "Sample silhouette by method",
               "K-Means versus Ward/average/complete linkage on the stratified samples.")
    add_figure(doc, "w3_compare_02_ari.png", "23", "Adjusted Rand Index",
               "Agreement between K-Means and Ward labels on the same samples (1 = identical partitions).")

    add_heading_1(doc, "10. Cluster Analysis and Interpretation")
    add_body(
        doc,
        "Labels below are descriptive summaries of computed statistics. They are not Netflix internal genre systems "
        "and they do not imply audience demand."
    )
    add_heading_2(doc, "10.1 Movie clusters")
    for cid in sorted(mp):
        p = mp[cid]
        add_heading_3(doc, f"Movie Cluster {cid}: {p['label']}")
        add_body(doc, f"Size: {p['n']:,} titles ({p['pct']}% of 5,377 movies).")
        add_bullet(doc, f"mean {p['mean_release_year']}, median {p['median_release_year']}.", "Release year: ")
        add_bullet(doc, f"mean {p['mean_duration_minutes']} min, median {p['median_duration_minutes']} min.", "Runtime: ")
        add_bullet(doc, f"mean {p['mean_content_age']} years, median {p['median_content_age']}.", "Content age at addition: ")
        add_bullet(doc, pct_join(p["audience_pct"], 4), "Audience mix: ")
        add_bullet(doc, pct_join(p["top_genres"], 5), "Frequent listed genres: ")
        add_bullet(doc, pct_join(p["country_group_pct"], 5), "Primary-country groups: ")
        add_bullet(doc, pct_join(p["top_ratings"], 4), "Common ratings: ")
        if "Classic" in p["label"]:
            add_body(doc, "This small vintage group differs from the newer catalogue in release year and addition lag. Possible use: flagging old titles for metadata review. Limitation: the group is small, so genre and rating percentages are less stable.")
        elif "Mid-catalogue" in p["label"]:
            add_body(doc, "This group has a longer addition lag and a teen-leaning US/India mix than the contemporary sets. Possible use: describing licensed back-catalogue features. Limitation: addition lag is not the same as cinematic importance or quality.")
        elif "Short-form" in p["label"]:
            add_body(doc, "This group is shorter and US-heavy, with documentary, stand-up, and family tags. Possible use: reviewing specials and shorter titles. Limitation: family titles and adult stand-up share a cluster, so it is also driven by length.")
        else:
            add_body(doc, "This is the largest contemporary international feature group, with dense genre tagging. Possible use: sampling non-US catalogue records for metadata QA. Limitation: ‘International Movies’ is a catalogue genre tag, not a verified production nationality.")

    add_heading_2(doc, "10.2 TV Show clusters")
    for cid in sorted(tp):
        p = tp[cid]
        add_heading_3(doc, f"TV Cluster {cid}: {p['label']}")
        add_body(doc, f"Size: {p['n']:,} titles ({p['pct']}% of 2,410 TV Shows).")
        add_bullet(doc, f"mean {p['mean_release_year']}, median {p['median_release_year']}.", "Release year: ")
        add_bullet(doc, f"mean {p['mean_seasons']} seasons, median {p['median_seasons']}.", "Season count: ")
        add_bullet(doc, f"mean {p['mean_content_age']} years, median {p['median_content_age']}.", "Content age at addition: ")
        add_bullet(doc, pct_join(p["audience_pct"], 4), "Audience mix: ")
        add_bullet(doc, pct_join(p["top_genres"], 5), "Frequent listed genres: ")
        add_bullet(doc, pct_join(p["country_group_pct"], 5), "Primary-country groups: ")
        add_bullet(doc, pct_join(p["top_ratings"], 4), "Common ratings: ")
        if "Legacy" in p["label"]:
            add_body(doc, "This small legacy group has earlier release years than the contemporary majority. Possible use: flagging older TV records for metadata review. Limitation: the group is small, so percentages are uncertain and do not imply programming strategy.")
        elif "Multi-season" in p["label"]:
            add_body(doc, "This contemporary, US-leaning group has a much higher season count than the mostly single-season groups. Possible use: studying long-running records in the catalogue. Limitation: season count is metadata, not evidence of current renewal.")
        elif "Older acquired" in p["label"]:
            add_body(doc, "This group has a longer addition lag than the contemporary majority. Possible use: reviewing older acquired series metadata. Limitation: addition lag does not imply a series is long-running.")
        else:
            add_body(doc, "This contemporary group is mostly single-season and carries the listed genres shown above. Possible use: describing common catalogue metadata patterns. Limitation: multiple genres remain mixed within the group.")

    add_heading_1(doc, "11. Model Evaluation and Comparison")
    add_body(
        doc,
        "Metrics are specific to these matrices. Hierarchical numbers use samples; K-Means numbers below use the full type subsets unless marked otherwise."
    )
    create_styled_table(doc, ["Experiment", "K", "Silhouette", "Davies–Bouldin", "Calinski–Harabasz", "Notes"], [
        ["K-Means movies full", "4", fmt(movies["final"]["silhouette"]), fmt(movies["final"]["davies_bouldin"]),
         fmt(movies["final"]["calinski_harabasz"], 1), "Primary movie model"],
        ["K-Means TV full", "4", fmt(tv["final"]["silhouette"]), fmt(tv["final"]["davies_bouldin"]),
         fmt(tv["final"]["calinski_harabasz"], 1), "Primary TV model"],
        ["K-Means movies sample", "4", fmt(movies["hierarchical"]["kmeans_sample_silhouette"]), "—", "—", "Same 250 rows as dendrogram"],
        ["Ward movies sample", "4", fmt(movies["hierarchical"]["ward_silhouette"]),
         fmt(movies["hierarchical"]["linkage_metrics"]["ward"]["davies_bouldin"]),
         fmt(movies["hierarchical"]["linkage_metrics"]["ward"]["calinski_harabasz"], 1),
         f"ARI vs K-Means {fmt(cmp_['movie_hier_ari'])}"],
        ["K-Means TV sample", "4", fmt(tv["hierarchical"]["kmeans_sample_silhouette"]), "—", "—", "Same 250 rows as dendrogram"],
        ["Ward TV sample", "4", fmt(tv["hierarchical"]["ward_silhouette"]),
         fmt(tv["hierarchical"]["linkage_metrics"]["ward"]["davies_bouldin"]),
         fmt(tv["hierarchical"]["linkage_metrics"]["ward"]["calinski_harabasz"], 1),
         f"ARI vs K-Means {fmt(cmp_['tv_hier_ari'])}"],
    ], col_widths=[1.7, 0.5, 1.0, 1.2, 1.3, 1.5])
    add_body(
        doc,
        "K-Means and Ward disagree on a substantial fraction of sampled titles (ARI ≈ 0.45–0.50) because Ward sees "
        "only the sample, uses a different merge criterion, and because the feature space is mixed binary/numeric. "
        "Average linkage’s higher sample silhouette does not make it a better catalogue taxonomy; it often rewards "
        "chaining small outliers. Computationally, K-Means on 5,377 rows is cheap; a full Ward dendrogram is the bottleneck, which is why sampling was used."
    )

    add_heading_1(doc, "12. Key Findings")
    add_bullet(doc, "The strongest geometric split in both types is recency / addition lag (K=2 silhouette peak).", "1. ")
    add_bullet(doc, "At K=4, movies separate into international features, shorter US docs/stand-up/family titles, mid-catalogue licensed features, and vintage cinema.", "2. ")
    add_bullet(doc, "At K=4, TV Shows separate into two contemporary mostly single-season groups with different genre/country mixes, a multi-season group, and a small legacy group.", "3. ")
    add_bullet(doc, "Genre multi-hot flags matter but do not produce pure single-genre clusters — titles are multi-labelled.", "4. ")
    add_bullet(doc, "PCA 2D views retain about 64% (movies) and 76% (TV) of variance; they are illustrations, not the clustering space.", "5. ")

    add_heading_1(doc, "13. Challenges and Solutions")
    add_bullet(doc, "Mixing minutes and seasons would manufacture false neighbours. Solution: split by type.", "Incommensurate duration: ")
    add_bullet(doc, "Silhouette favoured an uninformative majority/minority split. Solution: combine elbow, indices, sizes, and profiles; report K=2 and K=4.", "Metric conflict: ")
    add_bullet(doc, "42 genre tags would sparse-ify K-Means. Solution: frequent multi-hot lists only.", "Sparsity: ")
    add_bullet(doc, "Full dendrograms are unreadable. Solution: n=250 stratified sample with an explicit non-generalisation statement.", "Hierarchy scale: ")
    add_bullet(doc, "Release year and content age are related but measure original release timing and catalogue addition lag. Solution: document both as distinct catalogue dimensions and scale them robustly.", "Timing features: ")

    add_heading_1(doc, "14. Limitations and Future Scope")
    add_body(
        doc,
        "The table is catalogue metadata through early 2021, not a living service graph. Encoding and scaling choices "
        "change distances; another team using TF-IDF on descriptions or Gower mixed-type distance would obtain different "
        "groups. Cluster count remains uncertain. High-dimensional binary features weaken Euclidean intuition. "
        "Future work could include Gaussian mixtures, HDBSCAN, clustering on a Gower matrix, topic models on descriptions, "
        "or stability via bootstrap adjusted Rand indices — still without claiming recommender insight."
    )

    add_heading_1(doc, "15. Conclusion")
    add_body(
        doc,
        "Week 3 delivered a reproducible unsupervised analysis of the Netflix catalogue: type-separated feature matrices, "
        "Robust scaling, K-Means with a documented K search, Ward hierarchical clustering on qualified samples, PCA and "
        "profile visualisations, and interpretation that stays inside the metadata. The skills demonstrated are mixed-type "
        "preprocessing for distance methods, multi-criteria model selection, and cautious cluster storytelling."
    )

    add_heading_1(doc, "16. References")
    add_bullet(doc, "Shivam Bansal. Netflix Movies and TV Shows. Kaggle. https://www.kaggle.com/datasets/shivamb/netflix-shows", "1. ")
    add_bullet(doc, "Pedregosa et al. (2011). Scikit-learn: Machine Learning in Python. JMLR 12:2825–2830. https://scikit-learn.org/stable/modules/clustering.html", "2. ")
    add_bullet(doc, "sklearn.cluster.KMeans and sklearn.metrics (silhouette, Davies–Bouldin, Calinski–Harabasz) documentation.", "3. ")
    add_bullet(doc, "scipy.cluster.hierarchy linkage and dendrogram documentation.", "4. ")
    add_bullet(doc, "Jain, A. K. (2010). Data clustering: 50 years beyond K-means. Pattern Recognition Letters.", "5. ")
    add_bullet(doc, "Rousseeuw, P. J. (1987). Silhouettes: a graphical aid to the interpretation and validation of cluster analysis.", "6. ")

    add_heading_1(doc, "17. Appendix")
    add_heading_2(doc, "17.1 Movie K-search table")
    create_styled_table(
        doc,
        ["K", "Inertia", "Silhouette", "Davies–Bouldin", "Calinski–Harabasz", "Min n", "Max n"],
        [[str(r["k"]), fmt(r["inertia"], 1), fmt(r["silhouette"]), fmt(r["davies_bouldin"]),
          fmt(r["calinski_harabasz"], 1), str(int(r["min_cluster_size"])), str(int(r["max_cluster_size"]))]
         for r in movies["k_metrics"]],
        col_widths=[0.6, 1.1, 1.1, 1.3, 1.5, 0.8, 0.8],
    )
    add_heading_2(doc, "17.2 TV K-search table")
    create_styled_table(
        doc,
        ["K", "Inertia", "Silhouette", "Davies–Bouldin", "Calinski–Harabasz", "Min n", "Max n"],
        [[str(r["k"]), fmt(r["inertia"], 1), fmt(r["silhouette"]), fmt(r["davies_bouldin"]),
          fmt(r["calinski_harabasz"], 1), str(int(r["min_cluster_size"])), str(int(r["max_cluster_size"]))]
         for r in tv["k_metrics"]],
        col_widths=[0.6, 1.1, 1.1, 1.3, 1.5, 0.8, 0.8],
    )
    add_heading_2(doc, "17.3 Feature lists")
    add_body(doc, "Movie features: " + ", ".join(movies["feature_names"]))
    add_body(doc, "TV features: " + ", ".join(tv["feature_names"]))

    add_heading_1(doc, "18. Validation Checklist")
    create_styled_table(doc, ["Check", "Status"], [
        ["Week 3 notebook executed from a fresh Jupyter kernel", "Completed"],
        ["Clustering script executed (src/clustering.py)", "Completed"],
        ["Movies and TV Shows clustered separately", "Completed"],
        ["K-Means K=2…10 evaluated with multiple metrics", "Completed"],
        ["Elbow, silhouette, DB, CH charts saved", "Completed"],
        ["PCA and dendrogram figures saved under visualizations/week3/", "Completed"],
        ["Cluster labels backed by computed profiles", "Completed"],
        ["Hierarchical comparison documented as sample-based", "Completed"],
        ["Week 1 and Week 2 reports/notebooks left in place", "Completed"],
        ["Identifiers excluded from the feature matrix", "Completed"],
        ["No credentials committed", "Completed"],
        ["DOCX generated from the same metrics JSON as the pipeline", "Completed"],
        ["DOCX rendered for page-by-page visual QA", "Pending: LibreOffice soffice.exe unavailable"],
    ], col_widths=[4.6, 1.8])

    os.makedirs(os.path.dirname(REPORT_PATH), exist_ok=True)
    doc.save(REPORT_PATH)
    print(f"Report saved: {REPORT_PATH}")


if __name__ == "__main__":
    generate_report()
